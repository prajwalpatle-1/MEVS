"""Render generated summary Markdown as Word or PDF documents."""
from __future__ import annotations

import os
import re
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer


_IMAGE_PATTERN = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
_HEADING_PATTERN = re.compile(r"^(#{1,3})\s+(.*)$")
_BULLET_PATTERN = re.compile(r"^[-*]\s+(.*)$")


def _resolve_local_image(image_reference: str, output_dir: str) -> Path | None:
    """Resolve generated image links while preventing paths outside output_dir."""
    root = Path(output_dir).resolve()
    normalized = image_reference.replace("\\", "/")
    image_path = (root / normalized).resolve()
    try:
        image_path.relative_to(root)
    except ValueError:
        return None
    return image_path if image_path.is_file() else None


def _clean_inline_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return re.sub(r"[*_`]", "", text)


def _markdown_parts(markdown: str, output_dir: str):
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        image_match = _IMAGE_PATTERN.fullmatch(stripped)
        if image_match:
            image_path = _resolve_local_image(image_match.group(1), output_dir)
            if image_path:
                yield "image", image_path
            continue
        heading_match = _HEADING_PATTERN.match(stripped)
        if heading_match:
            yield f"heading{len(heading_match.group(1))}", _clean_inline_markdown(
                heading_match.group(2)
            )
            continue
        bullet_match = _BULLET_PATTERN.match(stripped)
        if bullet_match:
            yield "bullet", _clean_inline_markdown(bullet_match.group(1))
            continue
        yield "paragraph", _clean_inline_markdown(stripped)


def create_docx(markdown: str, output_dir: str) -> bytes:
    """Create a DOCX document from summary Markdown and local keyframes."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(42, 53, 64)

    for kind, content in _markdown_parts(markdown, output_dir):
        if kind.startswith("heading"):
            document.add_heading(content, level=int(kind[-1]))
        elif kind == "bullet":
            document.add_paragraph(content, style="List Bullet")
        elif kind == "image":
            document.add_picture(str(content), width=Inches(6.2))
        else:
            document.add_paragraph(content)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _register_pdf_fonts() -> None:
    font_dir = Path(os.path.dirname(__import__("reportlab").__file__)) / "fonts"
    pdfmetrics.registerFont(TTFont("MEVS", str(font_dir / "Vera.ttf")))
    pdfmetrics.registerFont(TTFont("MEVS-Bold", str(font_dir / "VeraBd.ttf")))


def create_pdf(markdown: str, output_dir: str) -> bytes:
    """Create a PDF document from summary Markdown and local keyframes."""
    _register_pdf_fonts()
    base_styles = getSampleStyleSheet()
    styles = {
        "heading1": ParagraphStyle(
            "MEVSHeading1", parent=base_styles["Heading1"], fontName="MEVS-Bold",
            fontSize=20, leading=24, textColor=colors.HexColor("#12304A"),
            spaceBefore=12, spaceAfter=8,
        ),
        "heading2": ParagraphStyle(
            "MEVSHeading2", parent=base_styles["Heading2"], fontName="MEVS-Bold",
            fontSize=15, leading=19, textColor=colors.HexColor("#087E8B"),
            spaceBefore=10, spaceAfter=6,
        ),
        "heading3": ParagraphStyle(
            "MEVSHeading3", parent=base_styles["Heading3"], fontName="MEVS-Bold",
            fontSize=12, leading=15, textColor=colors.HexColor("#12304A"),
            spaceBefore=8, spaceAfter=4,
        ),
        "paragraph": ParagraphStyle(
            "MEVSBody", parent=base_styles["BodyText"], fontName="MEVS",
            fontSize=10, leading=14, textColor=colors.HexColor("#2A3540"),
            spaceAfter=6,
        ),
        "bullet": ParagraphStyle(
            "MEVSBullet", parent=base_styles["BodyText"], fontName="MEVS",
            fontSize=10, leading=14, leftIndent=16, firstLineIndent=-10,
            textColor=colors.HexColor("#2A3540"), spaceAfter=4,
        ),
    }
    story = []
    for kind, content in _markdown_parts(markdown, output_dir):
        if kind == "image":
            story.append(Spacer(1, 6))
            story.append(
                Image(str(content), width=6.3 * inch, height=4.2 * inch,
                      kind="proportional", hAlign="CENTER")
            )
            story.append(Spacer(1, 10))
        elif kind == "bullet":
            story.append(
                Paragraph("&#8226;&nbsp;" + escape(content), styles["bullet"])
            )
        else:
            story.append(Paragraph(escape(content), styles[kind]))

    stream = BytesIO()
    pdf = SimpleDocTemplate(
        stream,
        pagesize=A4,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title="MEVS Video Summary",
        author="MEVS",
    )
    pdf.build(story)
    return stream.getvalue()